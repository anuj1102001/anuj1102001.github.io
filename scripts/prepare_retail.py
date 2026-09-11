"""Create monthly public-data aggregates for the portfolio explorer.
Source: Chen, D. (2015). Online Retail. UCI Machine Learning Repository.
https://doi.org/10.24432/C5BW33 — CC BY 4.0.
Usage: python scripts/prepare_retail.py /path/to/online-retail.zip
Requires openpyxl. No customer-level records are published.
"""
import sys,zipfile,json,io
from pathlib import Path
from collections import defaultdict
from decimal import Decimal
from openpyxl import load_workbook
archive=zipfile.ZipFile(sys.argv[1])
book=load_workbook(io.BytesIO(archive.read('Online Retail.xlsx')),read_only=True,data_only=True)
rows=book.active.iter_rows(values_only=True);next(rows)
months=[f'2011-{i:02}' for i in range(1,12)]
selected=['All markets','United Kingdom','Germany','France','Netherlands','EIRE']
agg={c:{m:{'revenue':Decimal(0),'invoices':set(),'rows':0} for m in months} for c in selected}
seen=set();raw=0;used=0
for row in rows:
 raw+=1
 invoice,stock,description,quantity,date,price,customer,country=row
 if row in seen: continue
 seen.add(row)
 if not date or date.year!=2011 or date.month>11:continue
 if str(invoice).upper().startswith('C') or not quantity or quantity<=0 or not price or price<=0:continue
 month=date.strftime('%Y-%m');revenue=Decimal(str(quantity))*Decimal(str(price));used+=1
 for c in ['All markets']+([country] if country in selected and country!='All markets' else []):
  a=agg[c][month];a['revenue']+=revenue;a['invoices'].add(str(invoice));a['rows']+=1
out={'source':'Chen, D. (2015). Online Retail. UCI Machine Learning Repository.', 'sourceUrl':'https://doi.org/10.24432/C5BW33','license':'CC BY 4.0','period':'January–November 2011','rawRows':raw,'retainedRows':used,'method':'Exact duplicate rows removed; positive quantity and price only; cancellation invoices excluded. All line items retained, including charges. Only complete months January–November 2011 shown. Orders are distinct retained invoice numbers.','months':months,'markets':{c:[{'month':m,'revenue':float(a['revenue'].quantize(Decimal('.01'))),'orders':len(a['invoices']),'rows':a['rows']} for m,a in data.items()] for c,data in agg.items()}}
Path('retail-data.json').write_text(json.dumps(out,separators=(',',':')))
print(json.dumps({'raw_rows':raw,'retained_rows':used,'all_market_revenue':round(sum(x['revenue'] for x in out['markets']['All markets']),2),'months':len(months)}))
